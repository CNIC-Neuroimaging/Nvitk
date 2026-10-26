# ─────────────────────────────────────────────────────────────────────────
# VENDORED FROM nvitk — DO NOT EDIT.
# Source: src/nvitk/measure/morpho/export_utils/report.py
# Regenerate: python MouseTOFMorphometricsLib/vendor_sync.py
# The only change from upstream is the root package rename nvitk -> nvitk_vendor.
# ─────────────────────────────────────────────────────────────────────────
"""The stage-7 morphometrics report: one row per vessel with the chosen features and the
high-confidence lesions, ``morphometrics_report.csv`` and the first sheet (``Report``) of
``case_metrics_donut_tree.xlsx``, and the per-vessel merge it is built from.

:func:`vessel_table` joins ``vessel_summary.csv`` with the earlier measures: cow_morpho's path
metrics (``00_Path_Summary``, aggregated per vessel by :func:`aggregate_summary_by_vessel`), the
label volumetry and the literature tortuosity metrics. A name the vessel summary already uses for
something else gets ``_legacy``. The GUI's results window shows the same table
(:mod:`nvitk_vendor.gui.viz.morpho_results`). Everything here is pandas only, so the pipeline (stage 7, on
SGE) and the GUI share it.

:func:`morphometrics_report` keeps :data:`REPORT_FEATURES` and summarises the lesions per vessel
from **high-confidence** lesions only (moderate and low ones are left out of the report):
``has_stenosis`` / ``stenosis_diameter_percent_max``, ``has_enlargement`` /
``enlargement_diameter_percent_max`` and ``has_hypoplasia`` (an absent main vessel).
``radius_mean_mm`` is cow_morpho's mean cross-section radius (length-weighted over the vessel's
segments), not half of ``diameter_mean_mm``, which is the inscribed one.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

#: The per-vessel features of the report (and of the database), in this order.
REPORT_FEATURES: tuple[str, ...] = (
    "centerline_length_mm", "diameter_median_mm", "diameter_mean_mm", "volume_mm3", "surface_area_mm2",
    "radius_mean_mm", "tortuosity_dm", "curvature_mean_1_per_mm", "total_curvature_rad",
    "inflection_count_metric", "chord_azimuth_angle_deg", "diameter_mean_blanked_mm",
    "chord_polar_angle_deg", "polar_angle_accumulated_change_deg",
)
#: The lesion columns of the report: from high-confidence lesions only.
REPORT_LESION_COLUMNS: tuple[str, ...] = (
    "has_stenosis", "stenosis_diameter_percent_max",
    "has_enlargement", "enlargement_diameter_percent_max",
    "has_hypoplasia",
)
REPORT_IDENTITY: tuple[str, ...] = ("case_id", "label", "vessel_name", "side", "present")
REPORT_CSV = "morphometrics_report.csv"
REPORT_SHEET = "Report"

# Aggregation: length-weighted means for continuous metrics; max / sum for stenosis.
_AGG_MEAN_COLS: tuple[str, ...] = (
    "radius_mean_mm",
    "radius_std_mm",
    "radius_cv",
    "radius_p05_mm",
    "radius_p50_mm",
    "radius_p95_mm",
    "taper_slope_mm_per_mm",
    "tortuosity_dm",
    "curvature_mean_1_per_mm",
    "curvature_median_1_per_mm",
    "curvature_p95_1_per_mm",
    "curvature_mean",
)
_AGG_MAX_COLS: tuple[str, ...] = (
    "curvature_max_1_per_mm",
    "stenosis_percent_max",
    "degree_of_stenosis_pct",
    "stenosis_percent",
    "radius_max_enlarged_mm",
    "enlargement_percent_max",
)
_AGG_MIN_COLS: tuple[str, ...] = (
    "radius_min_stenotic_mm",
)
_AGG_SUM_COLS: tuple[str, ...] = (
    "stenosis_length_total_mm",
    "stenosis_segments_n",
    "enlargement_length_total_mm",
    "enlargement_segments_n",
)


def aggregate_summary_by_vessel(df: pd.DataFrame) -> pd.DataFrame:
    """Collapse multiple path / component rows into one row per vessel_name.

    Morphometrics can emit several paths (tree arms) or CCs for the same vessel
    id; the GUI table should show one measure set per vessel.
    """
    if df.empty or "vessel_name" not in df.columns:
        return df
    key = "vessel_name"
    if df[key].astype(str).nunique(dropna=False) == len(df):
        return df

    rows: list[dict[str, Any]] = []
    for name, g in df.groupby(key, dropna=False, sort=False):
        row: dict[str, Any] = {key: name}
        if "full_name" in g.columns:
            fn = g["full_name"].dropna().astype(str)
            row["full_name"] = fn.iloc[0] if len(fn) else str(name)
        length = pd.to_numeric(g.get("length_mm"), errors="coerce")
        length_sum = float(length.fillna(0.0).sum())
        row["length_mm"] = length_sum
        weights = length.fillna(0.0).to_numpy(dtype=np.float64)
        w_sum = float(weights.sum())

        def _weighted_mean(col: str) -> float:
            """Length-weighted mean of *col* over the group (plain nanmean if weights sum to ~0)."""
            if col not in g.columns:
                return float("nan")
            vals = pd.to_numeric(g[col], errors="coerce").to_numpy(dtype=np.float64)
            if w_sum > 1e-9 and np.isfinite(vals).any():
                m = np.isfinite(vals) & (weights > 0)
                if m.any():
                    return float(np.average(vals[m], weights=weights[m]))
            return float(np.nanmean(vals)) if np.isfinite(vals).any() else float("nan")

        def _max(col: str) -> float:
            """Max finite value of *col* over the group, or NaN if none are finite."""
            if col not in g.columns:
                return float("nan")
            vals = pd.to_numeric(g[col], errors="coerce").to_numpy(dtype=np.float64)
            return float(np.nanmax(vals)) if np.isfinite(vals).any() else float("nan")

        def _min(col: str) -> float:
            """Min finite value of *col* over the group, or NaN if none are finite."""
            if col not in g.columns:
                return float("nan")
            vals = pd.to_numeric(g[col], errors="coerce").to_numpy(dtype=np.float64)
            return float(np.nanmin(vals)) if np.isfinite(vals).any() else float("nan")

        def _sum(col: str) -> float:
            """Sum of finite values of *col* over the group, or NaN if none are finite."""
            if col not in g.columns:
                return float("nan")
            vals = pd.to_numeric(g[col], errors="coerce").to_numpy(dtype=np.float64)
            return float(np.nansum(vals)) if np.isfinite(vals).any() else float("nan")

        for col in _AGG_MEAN_COLS:
            if col in g.columns:
                row[col] = _weighted_mean(col)
        for col in _AGG_MAX_COLS:
            if col in g.columns:
                row[col] = _max(col)
        for col in _AGG_MIN_COLS:
            if col in g.columns:
                row[col] = _min(col)
        for col in _AGG_SUM_COLS:
            if col in g.columns:
                row[col] = _sum(col)
        row["n_paths"] = int(len(g))
        rows.append(row)
    return pd.DataFrame(rows)


#: cow_morpho path metrics (``00_Path_Summary``), aggregated per vessel by
#: :func:`aggregate_summary_by_vessel`.
LEGACY_PATH_COLUMNS: tuple[str, ...] = (
    "length_mm", "n_paths",
    "radius_mean_mm", "radius_std_mm", "radius_cv", "radius_p05_mm", "radius_p50_mm", "radius_p95_mm",
    "radius_min_stenotic_mm", "taper_slope_mm_per_mm", "tortuosity_dm",
    "curvature_mean_1_per_mm", "curvature_median_1_per_mm", "curvature_p95_1_per_mm", "curvature_max_1_per_mm",
    "stenosis_percent_max", "degree_of_stenosis_pct", "stenosis_length_total_mm", "stenosis_segments_n",
    "radius_max_enlarged_mm", "enlargement_percent_max", "enlargement_length_total_mm", "enlargement_segments_n",
)
#: Label volumetry (``volumetry.csv`` / ``06_Volumetry``).
VOLUMETRY_COLUMNS: tuple[str, ...] = (
    "n_voxels", "volume_mm3", "volume_ul", "mesh_volume_mm3", "surface_area_mm2", "skeleton_length_mm",
    "equivalent_radius_mm", "input_volume_mm3", "measured_volume_fraction",
)
#: Same quantity under the same name in ``vessel_summary.csv``: kept once.
_SAME_MEANING = frozenset({"volume_mm3", "surface_area_mm2", "mesh_volume_mm3"})
#: Literature tortuosity (``centerlines/tortuosity_metrics.csv``) -> how it is combined per vessel.
TORTUOSITY_COLUMNS: dict[str, str] = {
    "tortuosity_index": "weighted", "inflection_count_metric": "weighted",
    "sum_of_angles_metric_deg_per_mm": "weighted", "total_curvature_deg_per_mm": "weighted",
    "bend_count": "sum", "bend_angle_sum_deg": "sum", "bending_length_mm": "max",
}


def _read_csv(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path) if path.is_file() else pd.DataFrame()
    except Exception:  # noqa: BLE001 (an unreadable table is an absent one)
        return pd.DataFrame()


def _volumetry(stage7_dir: Path) -> pd.DataFrame:
    vol = _read_csv(stage7_dir / "volumetry.csv")
    if vol.empty:
        try:
            vol = pd.read_excel(stage7_dir / "case_metrics_donut_tree.xlsx", sheet_name="06_Volumetry")
        except Exception:  # noqa: BLE001
            return pd.DataFrame()
    if "label" in vol.columns:
        vol = vol[vol["label"].astype(str) != "TOTAL"]
    return vol


def _vessel_of(name: str, vessel_names: list[str]) -> str:
    """The vessel a centerline file name belongs to (``LPCA_P2_M2i`` -> ``LPCA_P2``)."""
    matches = [v for v in vessel_names if name == v or name.startswith(v + "_")]
    return max(matches, key=len) if matches else ""


def _tortuosity_by_vessel(stage7_dir: Path, vessel_names: list[str]) -> pd.DataFrame:
    tort = _read_csv(stage7_dir / "centerlines" / "tortuosity_metrics.csv")
    if tort.empty or "centerline" not in tort.columns:
        return pd.DataFrame()
    tort = tort.assign(vessel_name=[_vessel_of(str(n), vessel_names) for n in tort["centerline"]])
    tort = tort[tort["vessel_name"] != ""]
    rows = []
    for name, g in tort.groupby("vessel_name", sort=False):
        w = pd.to_numeric(g.get("length_mm"), errors="coerce").fillna(0.0).to_numpy(dtype=float)
        row: dict[str, Any] = {"vessel_name": name}
        for col, how in TORTUOSITY_COLUMNS.items():
            if col not in g.columns:
                continue
            v = pd.to_numeric(g[col], errors="coerce").to_numpy(dtype=float)
            ok = np.isfinite(v)
            if not ok.any():
                row[col] = np.nan
            elif how == "sum":
                row[col] = float(v[ok].sum())
            elif how == "max":
                row[col] = float(v[ok].max())
            else:
                row[col] = float(np.average(v[ok], weights=w[ok])) if w[ok].sum() > 0 else float(v[ok].mean())
        rows.append(row)
    return pd.DataFrame(rows)


def _merge_extra(base: pd.DataFrame, extra: pd.DataFrame, columns: Any) -> pd.DataFrame:
    """*base* + the *columns* of *extra* (joined on ``vessel_name``); a name *base* already uses for
    something else gets ``_legacy``, one meaning the same is kept once."""
    if extra.empty or "vessel_name" not in extra.columns:
        return base
    cols = [c for c in columns if c in extra.columns and not (c in _SAME_MEANING and c in base.columns)]
    if not cols:
        return base
    part = extra[["vessel_name", *cols]].drop_duplicates("vessel_name")
    part = part.rename(columns={c: f"{c}_legacy" for c in cols if c in base.columns})
    return base.merge(part, on="vessel_name", how="left")


def vessel_table(stage7_dir: Path | str) -> pd.DataFrame:
    """One row per vessel: ``vessel_summary.csv`` + the earlier path, volumetry and tortuosity
    measures (module docstring). Falls back to ``00_Path_Summary`` for older outputs."""
    d = Path(stage7_dir)
    base = _read_csv(d / "vessel_summary.csv")
    try:
        paths = pd.read_excel(d / "case_metrics_donut_tree.xlsx", sheet_name="00_Path_Summary")
        legacy = aggregate_summary_by_vessel(paths)
        if "n_paths" not in legacy.columns and "vessel_name" in legacy.columns:
            legacy = legacy.assign(n_paths=1)
    except Exception:  # noqa: BLE001
        legacy = pd.DataFrame()
    if base.empty:
        if legacy.empty:
            return pd.DataFrame()
        keep = [c for c in ("label", "vessel_name", "full_name", *LEGACY_PATH_COLUMNS) if c in legacy.columns]
        base = legacy[keep].drop_duplicates("vessel_name").assign(present=True)
        if "label" not in base.columns and {"label", "vessel_name"} <= set(paths.columns):
            label_of = paths.drop_duplicates("vessel_name").set_index("vessel_name")["label"]
            base.insert(0, "label", base["vessel_name"].map(label_of))
    else:
        base = _merge_extra(base, legacy, LEGACY_PATH_COLUMNS)
    base = _merge_extra(base, _volumetry(d), VOLUMETRY_COLUMNS)
    names = [str(v) for v in base.get("vessel_name", pd.Series(dtype=str)).dropna().unique()]
    return _merge_extra(base, _tortuosity_by_vessel(d, names), tuple(TORTUOSITY_COLUMNS))


def report_lesion_columns(lesions: pd.DataFrame, vessel_names: list[str]) -> pd.DataFrame:
    """Per vessel, from the **high-confidence** lesions: whether it has a stenosis / an enlargement /
    a hypoplasia, and the largest diameter % of its stenoses / enlargements (NaN when none)."""
    rows = []
    if lesions.empty or not {"vessel_name", "kind"} <= set(lesions.columns):
        high = pd.DataFrame(columns=["vessel_name", "kind", "diameter_percent"])
    else:
        high = lesions[lesions.get("confidence_class", pd.Series("", index=lesions.index)).astype(str) == "high"]
    for name in vessel_names:
        sub = high[high["vessel_name"].astype(str) == str(name)]
        row: dict[str, Any] = {"vessel_name": name}
        for kind in ("stenosis", "enlargement"):
            k = sub[sub["kind"].astype(str) == kind]
            pct = pd.to_numeric(k["diameter_percent"], errors="coerce") if "diameter_percent" in k else pd.Series(dtype=float)
            row[f"has_{kind}"] = bool(len(k))
            row[f"{kind}_diameter_percent_max"] = float(pct.max()) if pct.notna().any() else np.nan
        row["has_hypoplasia"] = bool((sub["kind"].astype(str) == "hypoplasia").any())
        rows.append(row)
    return pd.DataFrame(rows)


def _hypoplasia_from_topology(vessels: pd.DataFrame) -> pd.DataFrame:
    """Hypoplasia rows for an output written before they were (eICAB topology assumed)."""
    try:
        from nvitk_vendor.measure.morpho.export_utils.vessel_tables import hypoplasia_rows
        from nvitk_vendor.measure.morpho.topology_io import load_eicab_topology

        return pd.DataFrame(hypoplasia_rows(vessels, load_eicab_topology()))
    except Exception:  # noqa: BLE001 (a non-eICAB label map has no main vessels to check)
        return pd.DataFrame()


def morphometrics_report(stage7_dir: Path | str) -> pd.DataFrame:
    """One row per vessel of the run (absent ones with ``present = False``): identity,
    :data:`REPORT_FEATURES` and :data:`REPORT_LESION_COLUMNS`."""
    d = Path(stage7_dir)
    table = vessel_table(d)
    if table.empty:
        return pd.DataFrame(columns=[*REPORT_IDENTITY, *REPORT_FEATURES, *REPORT_LESION_COLUMNS])
    lesions = _read_csv(d / "lesions.csv")
    kinds = set(lesions["kind"].astype(str)) if "kind" in lesions.columns else set()
    if "hypoplasia" not in kinds:
        extra = _hypoplasia_from_topology(table)
        if not extra.empty:
            lesions = pd.concat([lesions, extra], ignore_index=True)
    names = [str(n) for n in table["vessel_name"]]
    out = table[[c for c in REPORT_IDENTITY if c in table.columns]].copy()
    if "present" not in out.columns:
        out["present"] = True
    for col in REPORT_FEATURES:
        out[col] = pd.to_numeric(table[col], errors="coerce") if col in table.columns else np.nan
    out = out.merge(report_lesion_columns(lesions, names), on="vessel_name", how="left")
    if "label" in out.columns:
        out = out.sort_values("label")
    return out.reset_index(drop=True)


def write_morphometrics_report(stage7_dir: Path | str) -> Path:
    """Write :func:`morphometrics_report` as ``morphometrics_report.csv`` and as the first sheet
    (``Report``) of ``case_metrics_donut_tree.xlsx``; returns the CSV path."""
    d = Path(stage7_dir)
    report = morphometrics_report(d)
    csv = d / REPORT_CSV
    report.to_csv(csv, index=False)
    xlsx = d / "case_metrics_donut_tree.xlsx"
    if xlsx.is_file():
        from openpyxl import load_workbook

        with pd.ExcelWriter(xlsx, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            report.to_excel(writer, sheet_name=REPORT_SHEET, index=False)
        wb = load_workbook(xlsx)
        wb.move_sheet(REPORT_SHEET, offset=-wb.sheetnames.index(REPORT_SHEET))
        wb.active = 0
        wb.save(xlsx)
    return csv


__all__ = [
    "LEGACY_PATH_COLUMNS",
    "REPORT_CSV",
    "REPORT_FEATURES",
    "REPORT_LESION_COLUMNS",
    "REPORT_SHEET",
    "TORTUOSITY_COLUMNS",
    "VOLUMETRY_COLUMNS",
    "aggregate_summary_by_vessel",
    "morphometrics_report",
    "report_lesion_columns",
    "vessel_table",
    "write_morphometrics_report",
]
